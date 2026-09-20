from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED = {".md", ".markdown", ".txt", ".docx", ".pdf", ".xlsx", ".xlsm"}


@dataclass
class LoadedDoc:
    path: Path
    rel_path: str
    title: str
    product: str
    doc_type: str
    text: str


def _rel(path: Path, vault: Path) -> str:
    return path.resolve().relative_to(vault.resolve()).as_posix()


def _meta_from_rel(rel: str) -> tuple[str, str]:
    parts = rel.split("/")
    doc_type = parts[0] if len(parts) > 1 else "root"
    stem = Path(parts[-1]).stem
    product = parts[1] if doc_type == "products" and len(parts) > 2 else stem
    if doc_type != "products":
        product = stem
    return product, doc_type


def _title_from_text(text: str, fallback: str) -> str:
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("#"):
            return s.lstrip("#").strip() or fallback
        return s[:80]
    return fallback


def load_md_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def load_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower()
        if style.startswith("heading"):
            m = re.search(r"(\d+)", style)
            level = int(m.group(1)) if m else 2
            lines.append("#" * min(level, 6) + " " + t)
        else:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    lines: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            lines.append(f"[page {i}]\n{text}")
    return "\n".join(lines)


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return str(v)
    return str(v).strip().replace("\n", " ")


def load_excel(path: Path, max_rows_per_sheet: int = 5000, max_cell_chars: int = 500) -> str:
    """Convert workbook to text with fill-down context for merged-style bid sheets."""
    from openpyxl import load_workbook

    wb = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: list[str] = []
    try:
        for sheet in wb.worksheets:
            title = sheet.title or "Sheet"
            parts.append(f"# {title}")
            rows_raw: list[list[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell_str(c) for c in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if not any(cells):
                    continue
                rows_raw.append(cells)
                if len(rows_raw) >= max_rows_per_sheet:
                    break

            if not rows_raw:
                parts.append("[empty sheet]")
                continue

            # detect header-like row (many short labels) among first few rows
            header: list[str] = []
            data_start = 0
            for i, cells in enumerate(rows_raw[:8]):
                nonempty = [c for c in cells if c]
                if len(nonempty) >= 3 and sum(1 for c in nonempty if len(c) <= 12) >= max(2, len(nonempty) // 2):
                    header = cells
                    data_start = i + 1
                    break

            if header:
                parts.append(" | ".join(header))
            prev_ctx: list[str] = []
            col_n = len(header) if header else (max((len(r) for r in rows_raw), default=0))

            for cells in rows_raw[data_start:]:
                # fill-down: keep last non-empty value for leading/context columns
                padded = list(cells) + [""] * max(0, col_n - len(cells))
                padded = padded[:col_n] if col_n else padded
                ctx = list(prev_ctx)
                for i, val in enumerate(padded):
                    if val.strip():
                        if i >= len(ctx):
                            ctx.extend([""] * (i + 1 - len(ctx)))
                        ctx[i] = val.strip()
                    elif i < len(ctx) and ctx[i] and i < max(0, col_n - 2):
                        # inherit empty middle columns (merged cells)
                        padded[i] = ctx[i]
                prev_ctx = ctx
                # join with context; cap very long cells so one row doesn't dominate a chunk
                out_cells = []
                for c in padded:
                    c = (c or "").strip()
                    if len(c) > max_cell_chars:
                        c = c[:max_cell_chars] + "…"
                    out_cells.append(c)
                line = " | ".join(out_cells).strip(" |")
                if line:
                    parts.append(line)

            if len(rows_raw) >= max_rows_per_sheet:
                parts.append(f"[sheet truncated at {max_rows_per_sheet} rows]")
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return "\n".join(parts)


def load_document(path: Path, vault: Path) -> LoadedDoc | None:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        return None
    rel = _rel(path, vault)
    product, doc_type = _meta_from_rel(rel)
    if ext in {".md", ".markdown", ".txt"}:
        text = load_md_txt(path)
    elif ext == ".docx":
        text = load_docx(path)
    elif ext in {".xlsx", ".xlsm"}:
        text = load_excel(path)
    else:
        text = load_pdf(path)
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return None
    title = _title_from_text(text, path.stem)
    return LoadedDoc(
        path=path,
        rel_path=rel,
        title=title,
        product=product,
        doc_type=doc_type,
        text=text,
    )


def iter_docs(vault: Path):
    if not vault.exists():
        return
    for path in sorted(vault.rglob("*")):
        if not path.is_file():
            continue
        if path.name.startswith("."):
            continue
        if path.name.startswith("~$"):
            continue
        if path.suffix.lower() not in SUPPORTED:
            continue
        yield path
